from django.db.models import Prefetch
from django.http import HttpResponse, JsonResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from rest_framework.views import APIView

from access_manager.models import CardLog
from access_manager.serializers.card_log import CardLogFilterParams
from access_manager.swagger.card_log_export import swagger_export_card_logs
from main.models import DevicePublicSpaces

EXPORT_LIMIT = 10_000


class ExportCardLogsExcelView(APIView):
    @swagger_export_card_logs()
    def post(self, request, *args, **kwargs):
        params = CardLogFilterParams.check(request.data)
        tenant_id = request.user.tenant_id

        logs = (
            CardLog.objects.select_related("staff", "guest", "device__room")
            .prefetch_related(
                Prefetch(
                    "device__device_public_spaces",
                    queryset=DevicePublicSpaces.objects.select_related("public_space"),
                )
            )
            .list(
                filters=params.get("filters", {}),
                device_ids=params.get("device_ids", []),
                room_ids=params.get("room_ids", []),
                public_space_ids=params.get("public_space_ids", []),
                sort_by=params.get("sort_by", ["-event_ts"]),
                user_id=params.get("user"),
                card_num=params.get("card_num"),
                tenant_id=tenant_id,
            )
        )

        total_count = logs.count()
        if total_count == 0:
            return JsonResponse({"detail": "No logs found for given filters."}, status=404)

        truncated = total_count > EXPORT_LIMIT
        export_qs = logs[:EXPORT_LIMIT] if truncated else logs

        wb = Workbook()
        ws = wb.active
        ws.title = "Card Logs"

        headers = [
            "Card Number",
            "Event Timestamp",
            "Access Group",
            "Device",
            "Spaces",
            "User Type",
            "User Name",
            "Created At",
        ]
        header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        header_font = Font(bold=True, size=14, color="000000")
        header_alignment = Alignment(horizontal="center", vertical="center")

        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        data_alignment = Alignment(horizontal="center", vertical="center")
        col_widths = [0] * len(headers)

        for log in export_qs.iterator(chunk_size=2000):
            if log.staff:
                user_type, user_name = "Staff", log.staff.get_name()
            elif log.guest:
                user_type, user_name = "Guest", log.guest.get_name()
            else:
                user_type, user_name = "", ""

            space_parts = []
            if log.device.room:
                space_parts.append(f"Room {log.device.room.number}")
            for dps in log.device.device_public_spaces.all():
                space_parts.append(dps.public_space.name)
            space_field = ", ".join(space_parts)

            row = [
                log.number,
                log.event_ts.strftime("%Y-%m-%d %H:%M:%S"),
                log.get_access_group_display(),
                log.device.name,
                space_field,
                user_type,
                user_name,
                log.created_at.strftime("%Y-%m-%d %H:%M:%S"),
            ]
            ws.append(row)

            for cell in ws[ws.max_row]:
                cell.alignment = data_alignment

            for i, val in enumerate(row):
                if val:
                    col_widths[i] = max(col_widths[i], len(str(val)))

        if truncated:
            ws.append([])  # blank separator row
            warning_text = (
                f"WARNING: Export was limited to {EXPORT_LIMIT:,} logs out of {total_count:,} total. "
                "Please narrow your filters to export the remaining data."
            )
            warning_row = [warning_text] + [""] * (len(headers) - 1)
            ws.append(warning_row)
            warning_row_num = ws.max_row
            ws.merge_cells(start_row=warning_row_num, start_column=1, end_row=warning_row_num, end_column=len(headers))
            warning_cell = ws.cell(row=warning_row_num, column=1)
            warning_cell.font = Font(bold=True, size=14, color="FF0000")
            warning_cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            warning_cell.alignment = Alignment(horizontal="center", vertical="center")

        for i, width in enumerate(col_widths, start=1):
            adjusted = max(width + 1, 20)
            ws.column_dimensions[get_column_letter(i)].width = adjusted

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="card_logs.xlsx"'
        wb.save(response)

        if truncated:
            response["X-Message"] = (
                f"Export limited to {EXPORT_LIMIT:,} of {total_count:,} records. "
                "Narrow your filters to export remaining data."
            )
        else:
            response["X-Message"] = "CardLogs file sent successfully"
        return response
