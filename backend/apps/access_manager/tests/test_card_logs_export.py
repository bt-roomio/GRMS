from io import BytesIO

from django.urls import reverse
from openpyxl import load_workbook

from core.tests.base import BaseTestCase


class ExportCardLogsExcelTest(BaseTestCase):
    fixtures = (
        "tenant_profile.yaml",
        "tenant.yaml",
        "roles_permissions.yaml",
        "users.yaml",
        "room.yaml",
        "guest.yaml",
        "public_space.yaml",
        "customer.yaml",
        "device_profile.yaml",
        "device.yaml",
        "group.yaml",
        "group_public_space.yaml",
        "group_room.yaml",
        "staff.yaml",
        "card.yaml",
        "card_slot.yaml",
        "guest_public_space.yaml",
        "guest_card.yaml",
        "staff_card.yaml",
        "need_sync_device.yaml",
        "card_log.yaml",
    )

    def setUp(self):
        self.client.credentials(HTTP_AUTHORIZATION=self.karina_token)
        self.url = reverse("access_manager:export-card-logs-excel")

    # ── helpers ──────────────────────────────────────────────────────────
    def _load_workbook(self, response):
        """Parse the response content into an openpyxl Workbook."""
        return load_workbook(filename=BytesIO(response.content))

    # ── success cases ────────────────────────────────────────────────────
    def test_export_card_logs_success(self):
        """POST with empty payload should export all logs and return xlsx."""
        payload = {}
        response = self.client.post(self.url, payload, format="json")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("card_logs.xlsx", response["Content-Disposition"])
        self.assertEqual(response["X-Message"], "CardLogs file sent successfully")

    def test_export_card_logs_excel_headers(self):
        """Exported file must contain the expected column headers."""
        payload = {}
        response = self.client.post(self.url, payload, format="json")
        wb = self._load_workbook(response)
        ws = wb.active

        expected_headers = [
            "Card Number",
            "Event Timestamp",
            "Access Group",
            "Device",
            "Spaces",
            "User Type",
            "User Name",
            "Created At",
        ]
        actual_headers = [cell.value for cell in ws[1]]
        self.assertEqual(actual_headers, expected_headers)

    def test_export_card_logs_row_count(self):
        """Number of data rows must match the number of fixture card-logs."""
        payload = {}
        response = self.client.post(self.url, payload, format="json")
        wb = self._load_workbook(response)
        ws = wb.active

        # 4 fixture logs + 1 header row = 5 total rows
        data_rows = ws.max_row - 1
        self.assertEqual(data_rows, 4)

    def test_export_card_logs_staff_user_type(self):
        """Log associated with a staff member should have User Type = 'Staff'."""
        payload = {}
        response = self.client.post(self.url, payload, format="json")
        wb = self._load_workbook(response)
        ws = wb.active

        # Header: "Card Number" is col 1, "User Type" is col 6
        staff_row = None
        for row in ws.iter_rows(min_row=2, values_only=False):
            # card_log fixture pk c2d3e4f5 has staff_id set, number "65 28 23 12"
            if row[0].value == "65 28 23 12":
                staff_row = row
                break

        self.assertIsNotNone(staff_row, "Staff card-log row not found in export")
        self.assertEqual(staff_row[5].value, "Staff")  # User Type column

    def test_export_card_logs_guest_user_type(self):
        """Log associated with a guest should have User Type = 'Guest'."""
        payload = {}
        response = self.client.post(self.url, payload, format="json")
        wb = self._load_workbook(response)
        ws = wb.active

        guest_row = None
        for row in ws.iter_rows(min_row=2, values_only=False):
            # card_log fixture pk c1d2e3f4 has guest_id set, number "12 23 34 45"
            if row[0].value == "12 23 34 45":
                guest_row = row
                break

        self.assertIsNotNone(guest_row, "Guest card-log row not found in export")
        self.assertEqual(guest_row[5].value, "Guest")  # User Type column

    # ── 404 – no matching logs ───────────────────────────────────────────
    def test_export_card_logs_no_data_returns_404(self):
        """When filters match no logs, the endpoint must return 404."""
        payload = {
            "filters": {
                "from_date": "2099-01-01 00:00:00",
                "to_date": "2099-12-31 23:59:59",
            }
        }
        response = self.client.post(self.url, payload, format="json")

        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.json()["detail"], "No logs found for given filters.")

    # ── filter by date range ─────────────────────────────────────────────
    def test_export_card_logs_with_date_filter(self):
        """Filtering by date range returns only matching logs."""
        # Fixture events are at 10:00, 11:00, 12:00, 13:00 on 2025-01-15
        payload = {
            "filters": {
                "from_date": "2025-01-15 10:30:00",
                "to_date": "2025-01-15 12:30:00",
            }
        }
        response = self.client.post(self.url, payload, format="json")

        self.assertEqual(response.status_code, 200)
        wb = self._load_workbook(response)
        ws = wb.active

        # Should include 11:00 and 12:00 logs → 2 data rows
        data_rows = ws.max_row - 1
        self.assertEqual(data_rows, 2)

    # ── filter by card number ────────────────────────────────────────────
    def test_export_card_logs_with_card_num_filter(self):
        """Filtering by card_num returns only matching logs."""
        # payload = {
        #     "card_num": "09 87 65 98",
        # }
        payload = {
            "room_ids": [],
            "public_space_ids": [],
            "user": "5b66af57-fb27-4c26-9986-b9994e644605",
            "card_num": "12 23 34 45",
            "device_ids": [],
            "filters": {"from_date": "1999-07-21 10:00:00", "to_date": "2030-07-26 23:59:59"},
            "sort_by": ["-event_ts"],
            "size": 100,
            "page": 1,
        }
        response = self.client.post(self.url, payload, format="json")

        self.assertEqual(response.status_code, 200)
        wb = self._load_workbook(response)
        ws = wb.active

        data_rows = ws.max_row - 1
        self.assertEqual(data_rows, 1)
        self.assertEqual(ws.cell(row=2, column=1).value, "12 23 34 45")

    # ── filter by device_ids ─────────────────────────────────────────────
    def test_export_card_logs_with_device_ids_filter(self):
        """Filtering by device_ids returns only logs for those devices."""
        payload = {
            # device b2c3d4e5... is used by only 1 log (card "09 87 65 98")
            "device_ids": ["b2c3d4e5-f6a7-8901-bcde-f12345678901"],
        }
        response = self.client.post(self.url, payload, format="json")

        self.assertEqual(response.status_code, 200)
        wb = self._load_workbook(response)
        ws = wb.active

        data_rows = ws.max_row - 1
        self.assertEqual(data_rows, 1)
        self.assertEqual(ws.cell(row=2, column=1).value, "09 87 65 98")

    # ── unauthenticated request ──────────────────────────────────────────
    def test_export_card_logs_unauthenticated(self):
        """Request without auth credentials should be rejected."""
        self.client.credentials()  # clear credentials
        payload = {}
        response = self.client.post(self.url, payload, format="json")

        self.assertIn(response.status_code, (401, 403))

    # ── worksheet title ──────────────────────────────────────────────────
    def test_export_card_logs_worksheet_title(self):
        """Active worksheet should be titled 'Card Logs'."""
        payload = {}
        response = self.client.post(self.url, payload, format="json")
        wb = self._load_workbook(response)
        ws = wb.active

        self.assertEqual(ws.title, "Card Logs")
