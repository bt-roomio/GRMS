from django.db.models import Q
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from drf_yasg.utils import swagger_auto_schema
from rest_framework.response import Response
from rest_framework.views import APIView

from core.utils.permission import check_perms
from main.models import Device
from main.utils.get_device_space import get_space
from shuttle.models import TsKv
from shuttle.serializers.ts_kv import TagLogsFilterParams, TsKvFilterParams, TsKvFilterPath
from shuttle.swagger.tag_logs_export import swagger_export_tag_logs
from shuttle.utils.get_non_null_field import get_non_null_column
from shuttle.utils.permissions import WhiteListOrIsAuthenticated


class TsKvListView(APIView):
    permission_classes = (WhiteListOrIsAuthenticated,)

    @swagger_auto_schema(tags=["Shuttle, TsKv"])
    @check_perms(["shuttle.view_tskv"])
    def get(self, request, **kwargs):
        path = TsKvFilterPath.check(kwargs)
        device = Device.objects.filter(
            Q(id=path.get("entity_id"))  # pyright: ignore
            | Q(Q(tenant_id=path.get("tenant_id")) & Q(room_id=path.get("room_id")))  # pyright: ignore
        ).first()
        if not device:
            return Response({"detail": "Not found device."}, 404)

        params = TsKvFilterParams.check(request.GET)
        queryset, _ = TsKv.objects.by_device(device).get_history(**params)  # pyright: ignore
        return Response(queryset)

class ExportTsKvExcelView(APIView):

    @swagger_export_tag_logs()
    def post(self, request, *args, **kwargs):
        params = TagLogsFilterParams.check(request.data)
        tenant_id = request.user.tenant_id
        queryset = TsKv.objects.by_tenant(tenant_id).tag_logs(
            entity=params.get("device"),
            keys=params.get("keys"),
            start_ts=params.get("start_ts"),
            all_tags=params.get("all_tags"),
            sort_by=params.get("sort_by", ["-ts"]),
        )

        if not queryset.exists():
            return Response({"detail": "Not found."}, 404)

        device = params.get("device")
        spaces_row = get_space(device.id)
        spaces = ", ".join(spaces_row) if spaces_row else "None"

        workbook = Workbook()
        worksheet = workbook.active

        # --- Title row ---
        device_info = f"Tag Logs Export of device: {device.name} in spaces: {spaces}"
        worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
        title_cell = worksheet.cell(row=1, column=1)
        title_cell.value = device_info
        title_cell.font = Font(bold=True, size=14, color="FFFFFF")  # White text
        title_cell.fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")  # Dark grey bg
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        # --- Headers ---
        headers = ["Timestamp", "Key Name", "Value", "Data Type"]

        header_font = Font(bold=True, size=14, color="000000")
        header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        for col_num, header in enumerate(headers, 1):
            cell = worksheet.cell(row=2, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        # --- Data rows ---
        row_num = 3
        data_alignment = Alignment(horizontal="center", vertical="center")

        for record in queryset:
            ts_kv_obj = TsKv.objects.filter(
                ts=record.get("ts"),
                key__key=record.get("key_name"),
                entity=device
            ).first()

            if ts_kv_obj:
                data_dict = {
                    'bool_v': ts_kv_obj.bool_v,
                    'str_v': ts_kv_obj.str_v,
                    'long_v': ts_kv_obj.long_v,
                    'dbl_v': ts_kv_obj.dbl_v,
                    'json_v': ts_kv_obj.json_v,
                }
                field_name, field_value = get_non_null_column(data_dict)
                timestamp = record.get("ts").strftime("%Y-%m-%d %H:%M:%S")
                value_str = str(field_value) if field_value is not None else ""
                data_type = field_name.replace('_v', '') if field_name else "null"

                for col_num, value in enumerate(
                    [timestamp, record.get("key_name", ""), value_str, data_type], 1
                ):
                    cell = worksheet.cell(row=row_num, column=col_num)
                    cell.value = value
                    cell.alignment = data_alignment

                row_num += 1

        # --- Auto-adjust column widths ---
        for col_cells in worksheet.columns:
            max_length = 0
            col_letter = get_column_letter(col_cells[0].column)
            for cell in col_cells:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            adjusted_width = max_length + 1  # tighter spacing
            if adjusted_width < 15:  # set a reasonable minimum width
                adjusted_width = 15
            worksheet.column_dimensions[col_letter].width = adjusted_width

        # --- Return response ---
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="tag_logs_export.xlsx"'
        workbook.save(response)
        return response