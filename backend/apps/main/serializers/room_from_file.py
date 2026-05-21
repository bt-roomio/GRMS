import json
import re

from django.db import transaction
from openpyxl import load_workbook

from rest_framework import serializers

from main.models import Device, Room, RoomType

EXPORT_FORMATS = {"json", "xlsx"}
EXPORT_COLUMNS = ["number", "floor", "block", "type", "devices", "door_lock_device"]

REQUIRED_COLUMNS = {"number", "floor", "block", "type"}
OPTIONAL_COLUMNS = {"label", "door_lock_device", "devices"}
ALLOWED_COLUMNS = REQUIRED_COLUMNS | OPTIONAL_COLUMNS


def _normalize_header(header):
    return re.sub(r"\s+", "_", header.strip()).lower()


def _cell_to_str(value):
    if value is None:
        return ""
    return str(value).strip()


def _extract_device_names(value):
    if isinstance(value, list):
        return [str(v).strip() for v in value if str(v).strip()]
    return [d.strip() for d in _cell_to_str(value).split(",") if d.strip()]


class RoomFromFileSerializer(serializers.Serializer):
    file = serializers.FileField()

    def validate_file(self, file):
        name = file.name.lower()
        if not (name.endswith(".json") or name.endswith(".xlsx")):
            raise serializers.ValidationError("Only .json and .xlsx files are supported.")
        return file

    def _parse_json(self, file):
        try:
            data = json.loads(file.read().decode("utf-8-sig"))
        except (json.JSONDecodeError, UnicodeDecodeError) as e:
            raise serializers.ValidationError(f"Invalid JSON file: {e}")

        if not isinstance(data, list):
            raise serializers.ValidationError("JSON file must contain a list of room objects.")

        if not data:
            return set(), []

        headers = set()
        rows = []
        for item in data:
            if not isinstance(item, dict):
                raise serializers.ValidationError("Each item in the JSON array must be an object.")
            normalized = {_normalize_header(k): v for k, v in item.items()}
            headers.update(normalized.keys())
            rows.append(normalized)

        return headers, rows

    def _parse_xlsx(self, file):
        wb = load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers_row = next(rows_iter, None)
        if not headers_row:
            wb.close()
            return set(), []
        headers = [_normalize_header(str(h)) if h else "" for h in headers_row]
        rows = []
        for row in rows_iter:
            row_dict = {headers[i]: row[i] for i in range(len(headers)) if headers[i]}
            rows.append(row_dict)
        wb.close()
        return set(h for h in headers if h), rows

    def _validate_device_uniqueness(self, rows):
        seen = {}
        errors = []
        for idx, row in enumerate(rows, start=2):
            all_macs = _extract_device_names(row.get("devices"))
            door_lock = _cell_to_str(row.get("door_lock_device"))
            if door_lock:
                all_macs.append(door_lock)

            for mac in all_macs:
                if mac in seen:
                    errors.append(
                        {"row": idx, "message": f"Device '{mac}' is duplicated (first used in row {seen[mac]})."}
                    )
                else:
                    seen[mac] = idx
        return errors

    def validate(self, attrs):
        file = attrs["file"]
        name = file.name.lower()

        if name.endswith(".json"):
            headers, rows = self._parse_json(file)
        else:
            headers, rows = self._parse_xlsx(file)

        missing = REQUIRED_COLUMNS - headers
        if missing:
            raise serializers.ValidationError(f"Missing required columns: {', '.join(sorted(missing))}")

        if not rows:
            raise serializers.ValidationError("File contains no data rows.")

        attrs["_parsed_rows"] = rows
        return attrs

    def _is_room_unchanged(self, existing_room, device_names, door_lock_name):
        current_device_names = set(existing_room.devices.values_list("name", flat=True))
        current_door_lock_name = existing_room.door_lock_device.name if existing_room.door_lock_device else ""
        return current_device_names == set(device_names) and current_door_lock_name == door_lock_name

    def _validate_devices(self, tenant, device_names, door_lock_name, idx, existing_room=None):
        errors = []
        devices = []
        door_lock = None

        for name in device_names:
            try:
                device = Device.objects.get(tenant=tenant, name=name, is_active=True)
            except Device.DoesNotExist:
                errors.append({"row": idx, "message": f"Device '{name}' not found."})
                continue

            if device.device_public_spaces.exists():
                errors.append(
                    {
                        "row": idx,
                        "message": f"Device '{name}' is assigned to a public space and cannot be used.",
                    }
                )
                continue
            devices.append(device)

        if door_lock_name:
            try:
                door_lock = Device.objects.get(tenant=tenant, name=door_lock_name, is_active=True)
            except Device.DoesNotExist:
                errors.append({"row": idx, "message": f"Door lock device '{door_lock_name}' not found."})
                return devices, None, errors

            if door_lock.device_public_spaces.exists():
                errors.append(
                    {
                        "row": idx,
                        "message": f"Door lock device '{door_lock_name}' is assigned to a public space and cannot be used.",
                    }
                )

        return devices, door_lock, errors

    def _build_defaults(self, tenant, row, door_lock):
        room_type_obj = None
        type_val = _cell_to_str(row.get("type"))
        if type_val:
            room_type_obj, _ = RoomType.objects.get_or_create(title=type_val, tenant=tenant)

        defaults = {
            "type": room_type_obj,
            "state": [Room.Available],
            "status": "OFF",
            "door_lock_device": door_lock,
        }

        label = _cell_to_str(row.get("label"))
        if label:
            defaults["label"] = label

        return defaults

    def _build_error_message(self, errors):
        room_errors = [e for e in errors if e.get("type") == "room"]
        device_errors = [e for e in errors if e.get("type") == "device"]

        if room_errors:
            rows = ", ".join(str(e["row"]) for e in room_errors[:5])
            return f"Room validation failed at row(s) {rows}: missing required fields."

        if device_errors:
            macs = []
            for e in device_errors:
                match = re.search(r"'([^']+)'", e["message"])
                if match and match.group(1) not in macs:
                    macs.append(match.group(1))
            return (
                f"Invalid or unavailable devices: {', '.join(macs[:10])}. Please check MAC addresses before importing."
            )

        return "Import validation failed. Please check the file and try again."

    def create(self, validated_data):
        tenant = self.context["tenant"]
        rows = validated_data["_parsed_rows"]

        uniqueness_errors = self._validate_device_uniqueness(rows)
        if uniqueness_errors:
            macs = []
            for e in uniqueness_errors:
                match = re.search(r"'([^']+)'", e["message"])
                if match and match.group(1) not in macs:
                    macs.append(match.group(1))
            return {"message": f"Duplicate devices in file: {', '.join(macs[:10])}"}

        errors = []
        prepared = []

        for idx, row in enumerate(rows, start=2):
            number = _cell_to_str(row.get("number"))
            floor = _cell_to_str(row.get("floor"))
            block = _cell_to_str(row.get("block"))
            type_val = _cell_to_str(row.get("type"))

            if not (number and floor and block and type_val):
                errors.append({"row": idx, "type": "room"})
                continue

            device_names = _extract_device_names(row.get("devices"))
            door_lock_name = _cell_to_str(row.get("door_lock_device"))

            existing_room = Room.objects.filter(
                number=number,
                floor=floor,
                block=block,
                tenant=tenant,
            ).first()

            if existing_room and self._is_room_unchanged(existing_room, device_names, door_lock_name):
                prepared.append({"action": "skip"})
                continue

            devices, door_lock, device_errors = self._validate_devices(
                tenant,
                device_names,
                door_lock_name,
                idx,
                existing_room,
            )
            if device_errors:
                errors.extend({"row": e["row"], "type": "device", "message": e["message"]} for e in device_errors)
                continue

            prepared.append(
                {
                    "action": "update" if existing_room else "create",
                    "row": row,
                    "existing_room": existing_room,
                    "devices": devices,
                    "door_lock": door_lock,
                    "number": number,
                    "floor": floor,
                    "block": block,
                }
            )

        if errors:
            return {"message": self._build_error_message(errors)}

        created_count = 0
        updated_count = 0
        skipped_count = 0

        with transaction.atomic():
            for entry in prepared:
                if entry["action"] == "skip":
                    skipped_count += 1
                    continue

                if entry["door_lock"]:
                    Room.objects.filter(
                        door_lock_device=entry["door_lock"],
                        tenant=tenant,
                    ).update(door_lock_device=None)

                defaults = self._build_defaults(tenant, entry["row"], entry["door_lock"])
                room, created = Room.objects.update_or_create(
                    number=entry["number"],
                    floor=entry["floor"],
                    block=entry["block"],
                    tenant=tenant,
                    defaults=defaults,
                )

                if created:
                    created_count += 1
                else:
                    entry["existing_room"].devices.update(room=None)
                    updated_count += 1

                for device in entry["devices"]:
                    device.room = room
                    device.save(update_fields=["room"])

        return {"success": True, "message": "Import completed successfully !"}


class RoomExportSerializer(serializers.Serializer):
    room_ids = serializers.ListField(child=serializers.UUIDField(), default=[], required=False)
    format = serializers.ChoiceField(choices=sorted(EXPORT_FORMATS))

    def validate_room_ids(self, room_ids):
        tenant = self.context["tenant"]
        base_qs = (
            Room.objects.filter(tenant=tenant)
            .select_related(
                "type",
                "door_lock_device",
            )
            .prefetch_related("devices")
        )

        if not room_ids:
            return base_qs

        rooms = base_qs.filter(id__in=room_ids)
        found_ids = set(rooms.values_list("id", flat=True))
        missing_ids = set(room_ids) - found_ids
        if missing_ids:
            raise serializers.ValidationError(f"Rooms not found: {', '.join(str(i) for i in missing_ids)}")

        return rooms

    def _room_to_dict(self, room):
        return {
            "number": room.number,
            "floor": room.floor,
            "block": room.block,
            "type": room.type.title if room.type else "",
            "devices": list(room.devices.values_list("name", flat=True)),
            "door_lock_device": room.door_lock_device.name if room.door_lock_device else "",
        }

    def export(self):
        rooms = self.validated_data["room_ids"]
        return [self._room_to_dict(room) for room in rooms]
