import json

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework.parsers import MultiPartParser
from rest_framework import status
from rest_framework.response import Response
from rest_framework.views import APIView

from core.utils.permission import check_perms
from main.serializers.room_from_file import EXPORT_COLUMNS, RoomExportSerializer, RoomFromFileSerializer
from main.swagger.room_from_file import swagger_room_export, swagger_room_from_file_import


class RoomFromFileListView(APIView):
    parser_classes = [MultiPartParser]

    @swagger_room_from_file_import()
    @check_perms(["main.add_roomfromconf"])
    def post(self, request):
        serializer = RoomFromFileSerializer(data=request.data, context={"tenant": request.user.tenant})
        serializer.is_valid(raise_exception=True)
        result = serializer.save()
        if result.get("success", False):
            return Response(result, status=status.HTTP_200_OK)
        return Response(result, status=status.HTTP_400_BAD_REQUEST)


class RoomExportView(APIView):

    @swagger_room_export()
    @check_perms(["main.view_room"])
    def post(self, request):
        serializer = RoomExportSerializer(data=request.data, context={"tenant": request.user.tenant})
        serializer.is_valid(raise_exception=True)
        rows = serializer.export()
        export_format = serializer.validated_data["format"]

        if export_format == "json":
            return self._build_json_response(rows)
        return self._build_xlsx_response(rows)

    def _build_json_response(self, rows):
        content = json.dumps(rows, indent=2, ensure_ascii=False)
        response = HttpResponse(content, content_type="application/json")
        response["Content-Disposition"] = 'attachment; filename="rooms.json"'
        return response

    def _build_xlsx_response(self, rows):
        wb = Workbook()
        ws = wb.active
        ws.title = "Rooms"

        header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        header_font = Font(bold=True, size=14, color="000000")
        header_alignment = Alignment(horizontal="center", vertical="center")
        data_alignment = Alignment(horizontal="center", vertical="center")

        headers = [col.replace("_", " ").title() for col in EXPORT_COLUMNS]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        col_widths = [0] * len(EXPORT_COLUMNS)

        for row in rows:
            values = [", ".join(row[col]) if col == "devices" else row[col] for col in EXPORT_COLUMNS]
            ws.append(values)

            for cell in ws[ws.max_row]:
                cell.alignment = data_alignment

            for i, val in enumerate(values):
                if val:
                    col_widths[i] = max(col_widths[i], len(str(val)))

        for i, width in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(width + 1, 20)

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="rooms.xlsx"'
        wb.save(response)
        return response
